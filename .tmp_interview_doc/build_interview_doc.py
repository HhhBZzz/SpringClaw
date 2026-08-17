from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


BASE = Path("/Users/hanbingzheng/springclaw/.tmp_interview_doc/base.docx")
OUT = Path("/Users/hanbingzheng/Desktop/韩秉政-Agent项目面试话术-融合版.docx")


def set_east_asia_font(run, font_name="Microsoft YaHei", size=None, bold=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def normalize_paragraph(paragraph, size=10.5):
    for run in paragraph.runs:
        set_east_asia_font(run, size=size)


def set_para_text(paragraph, text, size=10.5, bold=False):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_east_asia_font(run, size=size, bold=bold)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_east_asia_font(run, size=16 if level == 1 else 13, bold=True)
    return p


def add_body(doc, text, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_east_asia_font(run, size=size)
    return p


def add_bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_east_asia_font(run, size=size)
    return p


def add_numbered(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_east_asia_font(run, size=size)
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(9)
        if idx != len(lines) - 1:
            run.add_break()
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_section_title(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_east_asia_font(run, size=12, bold=True)
    return p


def replace_springclaw_resume_block(doc):
    replacements = {
        "项目定位：面向企业级 Agent 平台构建 Java Agent Runtime / Harness 基座，围绕 Agent 生命周期管理、任务编排、工具调度、记忆上下文、沙箱隔离、Trace 可观测性和质量评估做工程化封装。":
            "项目定位：面向企业级 Agent 产品构建 Java Agent Runtime / Harness 基座，围绕多入口对话、任务编排、工具治理、Skills 扩展、上下文与记忆、运行 Trace、质量评估和控制台做工程化封装。",
        "技术栈：Java 17、Spring Boot 3.5、Spring AI 1.1、MyBatis-Plus、MySQL、Redis/Redisson、Redis Vector Store、RabbitMQ、SSE、WebSocket、飞书 SDK、Vue3、Python Skill。":
            "技术栈：Java 17、Spring Boot 3.5、Spring AI 1.1、MyBatis-Plus、MySQL、Redis/Redisson、Redis Vector Store、RabbitMQ、SSE、WebSocket、Vue3、Python/Markdown Skills。",
        "1. 抽象 Provider/Model/Runtime 配置模型，支持 Claude/DeepSeek/Qwen 等多模型切换、流式/非流式输出、故障转移、本地兜底和执行模式选择，避免模型硬编码。":
            "1. 抽象 Provider/Model/Runtime 配置模型，支持 Claude/DeepSeek/Qwen 等多模型切换、流式/非流式输出、故障转移、执行模式选择和确定性模型控制，避免模型硬编码与意图漂移。",
        "2. 构建 Agent 执行状态机，将请求拆成意图识别、上下文组装、工具选择、执行、反思校验和最终回答整理，并用 requestId/runId 串联日志、Trace 和运行步骤。":
            "2. 构建 Agent 执行主链路，将请求拆成路由判断、Context Snapshot 上下文构建、模型决策、工具提案、权限/风险校验、执行、事件记录和终态记忆抽取，并用 requestId/runId 串联 Trace。",
        "3. 重构 Memory/Context 体系：MySQL 会话事件流承载短期可审计记忆，Redis Vector Store + Embedding 承载长期语义召回，MemoryFrame/ContextSnapshot 统一上下文注入。":
            "3. 重构 Memory/Context 体系：Redis 承载短期会话记忆，MySQL memory_record 承载长期权威记忆，Redis Vector Store + Embedding 作为语义检索增强层，MemoryFrame/ContextSnapshot 统一上下文注入。",
        "4. 封装 Tool Calling 治理链路：基于 Spring AI @Tool + AOP 统一工具注册、权限校验、限流、审计、风险分级和 ToolExecutionContext；对写入/命令等副作用操作生成审批 Proposal。":
            "4. 封装 Tool Calling 治理链路：模型只生成工具调用提案，后端负责工具注册、参数校验、权限校验、风险分级、审计 Trace；对写入/命令等副作用操作生成确认 Proposal，避免模型越权执行。",
        "5. 建设 Skill/Workflow 扩展机制：通过 SKILL.md 管理 Python/Builtin/Prompt 技能，支持技能扫描、运行统计和热扩展；结合 Workspace Guard 做工作区访问边界与副作用控制。":
            "5. 建设 Skills/Workflow 扩展机制：通过技能目录、Markdown Skill、脚本技能和 CapabilityRegistry 管理 Agent 能力包，支持技能扫描、策略配置、运行统计和热扩展，并接入权限与风险控制。",
        "6. 设计 Agent 运行态质量评估：按路由、工具、证据、反思、回答、成本、风险 7 个维度生成 quality score / level，并将 evaluation_json 落库用于 Trace 面板复盘。":
            "6. 设计运行态可观测与评估：记录路由、上下文、记忆注入、工具调用、模型用量、反思校验和 memory usage trace，用 runtime console 支持 run replay、记忆候选审核、provider harness 与红线评估。"
    }
    for p in doc.paragraphs:
        text = p.text.strip()
        if text in replacements:
            set_para_text(p, replacements[text])


def append_interview_script(doc):
    add_page_break(doc)
    add_heading(doc, "SpringClaw Agent 项目面试话术（融合当前实现版）", 1)
    add_body(doc, "说明：这部分用于面试口述，重点按照“我做的是什么、为什么做、整体架构、执行链路、负责模块、问题与解决、项目收获”的逻辑展开。英文术语第一次出现时附中文解释。")

    add_heading(doc, "1. 我做的是什么", 2)
    for text in [
        "我这个项目做的是一个基于 Java / Spring Boot 的 AI Agent 后端系统，我更愿意把它叫成一个 Agent 工作台，而不是简单聊天机器人。",
        "普通聊天机器人一般是用户发一句话，后端把这句话发给大模型，再把结果返回。但我这个项目想解决的是更接近真实 Agent 产品的问题：用户不只是问答，还可能让系统记住偏好、切换模型、调用工具、分析项目、执行任务、处理 webhook 事件，甚至在后台跑异步任务。",
        "所以它后端不能只是一个大模型 API 转发器，而是需要一套 Agent Runtime，也就是智能体运行时。它负责把用户的一句自然语言请求，变成一条可追踪、可校验、可执行、可回放的工程链路。",
        "从产品功能上看，它包含多模式聊天、模型切换、上下文理解、短期记忆、长期记忆、向量检索、工具调用、Skills 技能机制、任务系统、webhook 接入、运行轨迹追踪、记忆审核、模型调用统计和后台管理控制台。"
    ]:
        add_body(doc, text)

    add_heading(doc, "2. 为什么做", 2)
    for text in [
        "我做这个项目的出发点是：AI Agent 真正难的地方不在于能不能调用模型，而在于模型接入真实业务以后，怎么保证稳定、可控和可追踪。",
        "如果只是做 demo，controller 里调一下模型接口就能跑。但一旦要做成产品，会马上遇到上下文、记忆、工具安全、运行可观测、成本和延迟这些问题。",
        "比如用户连续对话时，系统到底该记住什么？如果把所有历史消息都塞进 prompt，短期看可以，长期会 token 爆炸，而且无关内容会干扰模型。",
        "再比如 Function Calling，也就是函数/工具调用，不是让模型完全替后端做业务决策。模型可以判断可能要调用哪个工具，但真正能不能执行，必须由后端根据权限、风险、状态和用户确认来决定。",
        "所以我做这个项目，本质上是想验证一套生产级 Agent 后端工程模式：让大模型负责理解和生成，让后端负责边界、状态、校验、权限、记忆和回放。"
    ]:
        add_body(doc, text)

    add_heading(doc, "3. 整体架构", 2)
    add_body(doc, "整体上我把 SpringClaw 分成七层。")
    for text in [
        "第一层是接入层。它支持普通聊天、流式聊天、异步聊天、webhook 和定时任务。不同入口进来后，后端都会生成 requestId 或 runId，让一次请求可追踪。",
        "第二层是路由层。这里我没有完全依赖大模型做意图识别，而是做了规则优先、模型兜底的混合路由。像“切换 DeepSeek 模型”这种确定性命令，优先用规则识别；普通偏好问题走轻量对话；涉及工作区、文件、脚本的请求，才进入对应工具或复杂 Agent 流程。",
        "第三层是上下文构建层。核心是 Context Snapshot，也就是上下文快照。每次模型调用前，系统会把当前问题、最近对话、Redis 短期记忆、MySQL 长期记忆、项目记忆、工具状态统一整理，做去重、筛选和 token 预算控制，再喂给模型。",
        "第四层是模型决策层。系统支持多个模型 provider，比如 DeepSeek、Claude 或其他模型。不同任务可以选择不同模型，比如普通问答、结构化抽取、反思校验可以使用不同 provider。",
        "第五层是工具和 Skills 层。工具偏底层能力，比如文件、工作区、脚本、系统状态；Skills 更像 Agent 的能力包，可以通过目录、Markdown 或脚本扩展。模型不是看到无限工具，而是后端根据 intent、用户权限、风险等级动态约束可用能力。",
        "第六层是记忆层。短期记忆用 Redis，长期记忆用 MySQL 的 memory_record，向量检索通过 Spring AI VectorStore 接入。MySQL 是权威记忆源，向量库是语义检索增强，不是唯一事实来源。",
        "第七层是观测和控制台。系统有运行时控制台，可以看模型状态、工具、技能、任务、最近运行、token 使用、记忆候选、记忆评估和 run trace。"
    ]:
        add_body(doc, text)
    add_code_block(doc, [
        "入口请求 → 路由判断 → 上下文构建 → 模型决策",
        "      → 工具/技能治理 → 结果持久化 → 记忆抽取 → 运行追踪和评估"
    ])

    add_heading(doc, "4. 核心执行链路：真实案例", 2)
    for text in [
        "我用一个真实验证过的案例来讲完整链路。用户先说：请长期记住：我以后做后端项目优先使用 Java、Spring Boot 和 MySQL。",
        "这句话进入聊天接口后，后端会生成 requestId。这个 requestId 会贯穿整个链路，用来记录 message_event、run trace、记忆抽取和前端展示。",
        "然后进入路由层。系统会判断这不是文件操作，也不是复杂项目分析，而是普通对话里带有长期记忆意图，所以不会进入重型 OPAR 流程。OPAR 可以理解成观察、计划、行动、反思的复杂 Agent 流程，它适合复杂任务，但不适合所有问题都走。",
        "接着进入上下文构建。系统会读取当前 session 的短期上下文、当前用户的长期记忆、项目记忆和工具状态，形成 Context Snapshot。这里不会把全部历史消息一股脑塞给模型，而是整理后再注入。",
        "用户得到回复后，真正关键的是对话结束后的终态异步处理。TerminalMemoryExtractionService 会读取这次 run 里的 message_event，让模型从对话中抽取稳定事实。它不是把原话直接存进去，而是抽象成一条长期记忆：用户做后端项目时，优先使用 Java、Spring Boot 和 MySQL。",
        "这条记忆会写入 MySQL 的 memory_record 表，带上 memory type、status、confidence、importance 和 evidence_refs。evidence_refs 是证据引用，表示这条记忆来自哪次 run、哪条 event，避免模型凭空制造长期记忆。",
        "后面用户换一个 session 再问：我做后端项目时优先喜欢用什么技术栈？这次请求进来后，系统仍然先路由。这个问题虽然包含“后端项目、技术栈”，但我修过路由，不让它误判成工作区分析，而是让它走轻量问答。",
        "然后 MemoryCoordinator 从 memory_record 里读取当前用户的 ACTIVE 长期记忆，发现有一条和后端技术栈相关的记忆，于是注入上下文。模型拿到的不只是用户当前问题，还有长期偏好，最终回答 Java + Spring Boot + MySQL。",
        "系统还会在事件里记录 memoryInjected=true 和 memoryReferencedInAnswer=true。前者表示这次确实注入了记忆，后者表示回答确实使用了这条记忆。这样我不只是知道有没有存，还知道这条记忆有没有真正发挥作用。"
    ]:
        add_body(doc, text)

    add_heading(doc, "5. Skills 机制", 2)
    for text in [
        "我在项目里对 Skills 的理解，不是简单多加几个工具，而是把 Agent 的能力模块化。",
        "如果所有能力都硬编码在 Agent 里，后面会越来越难维护。比如文件分析、工作区搜索、脚本执行、项目审查、知识源同步，这些能力的触发条件、参数、风险等级、权限要求都不一样。如果全部写在一个大 service 里，后期会非常混乱。",
        "所以我把 Skills 理解成 Agent 的能力包。一个 Skill 可以描述它能做什么、什么时候触发、需要什么参数、风险等级是什么、执行方式是什么。系统里有 Skill Catalog，也就是技能目录；有 Markdown Skill 导入；有脚本技能；也有 Skill Policy，也就是技能策略。",
        "用户输入进来后，路由层和 CapabilityRegistry 会判断这句话可能命中哪些能力。CapabilityRegistry 可以理解成能力注册表，里面有 workspace、file、script、system、weather、web、knowledge source 等能力。然后系统会根据用户角色、当前 channel、风险等级和工具权限，决定最终暴露哪些能力。",
        "比如用户说“帮我检查项目里 memory_record 相关实现”，这类请求可能命中 workspace-search 或 workspace-review 能力。如果只是读文件和分析代码，风险比较低；但如果用户说“帮我直接修改这个文件”，就涉及写操作，风险更高，需要进入工具调用提案并等待用户确认。",
        "所以 Skills 机制的核心价值是：Agent 的能力可以扩展，但扩展出来的能力仍然要被权限、风险和确认机制管理。它不是让模型看到所有能力后自由发挥，而是后端先筛一遍，把可用能力限制在安全范围里。"
    ]:
        add_body(doc, text)

    add_heading(doc, "6. Function Calling 和工具调用治理", 2)
    for text in [
        "我在项目里对 Function Calling 的理解，不是让模型替我执行后端逻辑，而是把模型的语言理解能力收进一个受控动作集合里。",
        "模型可以判断用户可能需要哪个能力，但真正执行时，必须由后端兜底。因为模型有不确定性，它可能工具名错、参数错，也可能把普通问题误判成工具调用。如果工具涉及文件修改、脚本执行、系统操作，风险就更大。",
        "所以我在 SpringClaw 里做的是工具提案机制。模型不是直接执行工具，而是先产生 tool proposal，也就是工具调用提案。后端再校验工具是否存在、用户是否有权限、参数是否合法、风险等级是什么、当前 run 状态是否允许继续。",
        "如果是低风险读操作，比如读取项目文件，系统可以直接执行。如果是高风险操作，比如改文件、执行脚本，就必须等待用户确认。用户确认以后，后端才会进入真正执行流程。",
        "这里的关键是：Function Calling 不能只有模型输出工具名这一层，它还应该有执行前约束、执行中治理和执行后留痕。执行前约束是能力注册、权限校验、参数校验和风险判断；执行中治理是确认机制、限流、状态机和幂等控制；执行后留痕是 message_event、tool audit、run trace 和执行结果记录。",
        "我之前遇到过一个真实问题：模型返回 ask_clarification，也就是澄清问题，但系统曾经把它误当成需要确认的 action。用户只是应该被问“你想怎么处理”，结果前端展示成类似确认执行的卡片。后来我把澄清问题和工具确认分成两条路径。这个问题让我意识到，Agent 产品里模型意图和产品动作必须严格区分。"
    ]:
        add_body(doc, text)

    add_heading(doc, "7. 记忆、RAG 和 Query Rewrite", 2)
    for text in [
        "我在项目里对记忆系统的理解，不是把聊天记录存起来，而是把过去对当前任务有用的信息，经过筛选、抽取、分层、检索后，在正确时机放进上下文。",
        "第一层是短期记忆。短期记忆主要解决当前 session 里的连续追问，比如用户说“继续刚才那个方案”。这类信息时效性强，所以我用 Redis 存，读写快，也适合做热上下文。",
        "第二层是长期记忆。长期记忆保存用户偏好、稳定事实和历史决策。比如用户偏好 Java 技术栈，或者用户不希望某类回答太抽象。这类信息需要可审计、可修改、可删除、可替换，所以我用 MySQL 的 memory_record 作为权威数据源。",
        "第三层是向量检索。项目里有 Spring AI VectorStore，可以接 Redis VectorStore。向量检索的作用是做语义召回增强，比如用户换一种说法问，系统也能找到相关内容。但我没有把向量库当成唯一事实来源，因为向量库更适合找相似内容，不适合管理长期事实的状态和版本。",
        "所以我的设计是：MySQL memory_record 负责事实权威，VectorStore 负责语义召回增强，Redis 负责短期上下文，Context Snapshot 负责最终注入模型。",
        "这其实是一种 Agent Memory RAG，中文可以叫智能体记忆型检索增强生成。它和传统文档 RAG 不完全一样。传统 RAG 更多是文档切片、embedding、topK 检索、塞进 prompt；我这里检索对象更多是用户偏好、项目约束、历史决策和反思经验。",
        "Query Rewrite，也就是查询改写，在项目里目前不是一个独立大模块，而是体现在路由归一化、记忆检索 query 构造和工具能力匹配上。比如“我做后端项目时优先喜欢用什么技术栈”会被归一成“查询当前用户的后端技术栈偏好”，而不是误判成代码项目分析。后续如果继续增强，我会把它抽成独立 QueryRewriteService，同时生成 keyword query、vector query 和 rule query，分别给关键词检索、向量检索和规则过滤使用。"
    ]:
        add_body(doc, text)

    add_heading(doc, "8. 任务系统、Webhook 和运行时控制台", 2)
    for text in [
        "除了聊天，我也做了任务和 webhook 入口。任务系统解决的是 Agent 主动执行的问题。用户可以创建任务、启用任务、禁用任务、手动运行任务，也可以让系统从自然语言里解析任务草稿。比如用户说“每天早上检查项目状态”，系统可以先生成任务草稿，让用户确认后再创建定时任务。",
        "Webhook 解决的是外部系统接入问题。外部系统可以把事件发给 SpringClaw，后续统一进入 Agent Runtime。这让 Agent 不只是网页聊天框，而是可以被用户消息、定时任务、外部事件共同触发的后端系统。",
        "运行时控制台对 Agent 产品非常重要。控制台可以看模型 provider、工具、技能、任务、最近 runs、模型使用、记忆候选、知识源、记忆评估和 run trace。比如用户反馈系统没有记住他，我可以看这次对话有没有触发记忆抽取、memory_record 有没有写入、状态是 CANDIDATE 还是 ACTIVE、下一次回答有没有注入记忆、回答有没有真正引用这条记忆。"
    ]:
        add_body(doc, text)

    add_heading(doc, "9. 我负责 / 实现了哪些模块", 2)
    for text in [
        "这个项目里，我主要负责整体 Agent Runtime 的设计和实现，包括路由、上下文、记忆、工具治理、Skills 扩展、运行追踪和测试验收。",
        "第一是路由机制，包括普通问答、模型控制、记忆问题、工作区任务、复杂任务的区分。我修复过真实体验问题，比如偏好问题不再误入重型项目分析链路。",
        "第二是上下文和记忆，包括 Redis 短期记忆、MySQL 长期记忆、Context Snapshot、MemoryCoordinator、语义抽取、证据引用、记忆候选 review 和记忆使用 trace。",
        "第三是工具调用治理，包括 tool proposal、确认、取消、执行、权限、风险等级和运行事件记录。",
        "第四是 Skills 机制，包括技能目录、Markdown 技能导入、脚本技能 reload、能力注册和技能策略。",
        "第五是任务系统和 runtime console，包括任务管理、运行记录、模型 provider 管理、记忆候选、记忆评估和 run trace。",
        "第六是测试体系，包括单元测试、架构测试、红线测试和真实接口 smoke test。因为 Agent 项目很容易看起来能跑但体验不对，所以我后来非常重视真实产品验证。"
    ]:
        add_body(doc, text)

    add_heading(doc, "10. 遇到的最大问题", 2)
    for text in [
        "我遇到的最大问题是：Agent 的链路太长，任何一环错了，用户看到的就是“这个 Agent 很笨”。",
        "最典型的是意图识别问题。用户说“切换 DeepSeek 模型”，早期可能被模型当成普通聊天。用户问“我做后端项目喜欢什么技术栈”，因为里面有“项目、技术栈”，系统可能误判成工作区分析，进入重型流程，结果又慢又不自然。",
        "第二个问题是记忆污染。早期如果把每轮对话都直接写入向量库，系统可能把临时表达当成长期偏好。比如用户说“我可能以后试试 Go”，这不一定代表他长期偏好 Go。如果直接存，就会污染后续回答。",
        "第三个问题是上下文断片。用户连续问问题时，如果短期记忆和长期记忆没有分层，模型可能不知道用户“刚才说的那个”到底指什么。",
        "第四个问题是工具执行风险。模型如果误判工具调用，或者参数生成错了，不能让它直接影响文件、系统或外部业务状态。",
        "第五个问题是 AI 辅助开发本身也会带来问题。这个项目我大量用了 Codex 辅助写代码、重构和补测试，但 Codex 有时会倾向于把架构做复杂，或者修了技术问题但没有解决产品体验问题。所以人必须一直把控目标。"
    ]:
        add_body(doc, text)

    add_heading(doc, "11. 怎么解决", 2)
    for text in [
        "我解决这些问题的思路是：把模型的不确定性限制在后端可治理的范围里。",
        "对意图识别，我做了轻重链路分离。普通问答、偏好记忆、模型切换走轻量链路，复杂任务才走 OPAR 或多步 Agent 流程。明确命令优先规则处理，不让模型猜。",
        "对上下文，我做了 Context Snapshot。每次请求前统一构建上下文，把短期记忆、长期记忆、项目记忆、工具状态整合起来，再筛选和控制预算。",
        "对记忆污染，我退役了“每轮对话直接写向量库”的生产路径，改成终态异步抽取稳定记忆。长期记忆进入 memory_record，有状态、有证据、有置信度，低置信进入候选 review。",
        "对工具风险，我做 proposal + confirmation。模型只能提出工具提案，高风险操作必须用户确认。后端还要做权限、风险、状态校验和审计。",
        "对运行不可观测，我做 trace 和 message_event。每次路由、记忆注入、工具调用、模型回答都会记录。这样出问题时能复盘，不是靠猜。",
        "对 AI 辅助开发，我建立了测试和产品验证习惯。Codex 生成代码以后，我会看 diff、跑测试、做真实接口 smoke test，再决定是否合并。"
    ]:
        add_body(doc, text)

    add_heading(doc, "12. Codex 在项目里的作用", 2)
    for text in [
        "这个项目里我大部分编码、重构、测试补齐和 PR 整理，确实大量使用了 Codex 辅助。",
        "但我不会把它描述成项目都是 AI 自动做的。更准确的说法是：我把 Codex 当成高级开发助手，用来提高开发效率。",
        "我主要让 Codex 做拆任务、生成初版代码、补测试、扫调用链、做重构、跑测试、分析失败原因、生成审查报告和整理 PR。",
        "但项目方向不是 Codex 决定的。比如长期记忆为什么用 MySQL 做权威源、为什么向量库只做检索增强、为什么高风险工具要确认、为什么普通问题不能进重型链路，这些是我根据产品体验和工程边界做的判断。",
        "所以面试时我会说：我大量使用 Codex 辅助开发，但我负责的是需求拆解、架构判断、边界控制、测试验收和最终合并。Codex 提升了实现效率，但不能替代工程判断。"
    ]:
        add_body(doc, text)

    add_heading(doc, "13. 项目收获", 2)
    for text in [
        "这个项目让我对 AI Agent 的理解发生了变化。一开始我以为 Agent 的核心是模型能力和工具数量。后来我发现，真正重要的是运行时治理。",
        "一个 Agent 如果要进入真实业务，必须有上下文管理，不能每次都孤立回答，也不能把所有历史都塞进去。",
        "它要有分层记忆。短期记忆、长期记忆、项目记忆、向量检索要分清楚。",
        "它要有工具治理。模型可以提出动作，但执行权必须在后端。",
        "它要有可观测性。每次运行要能 trace、能 replay、能定位问题。",
        "它还要有产品验证。架构再漂亮，如果用户问一个基础上下文问题都答不好，那就是没做好。",
        "所以这个项目对我最大的价值是：我不再把 AI 应用理解成调大模型接口，而是理解成把大模型能力工程化。也就是用 Java 后端的状态机、数据库、缓存、队列、权限、审计和测试，把大模型的不确定性控制在一个可接受的产品范围内。"
    ]:
        add_body(doc, text)

    add_heading(doc, "14. 2 分钟压缩版", 2)
    for text in [
        "我这个项目是一个 Java / Spring Boot 的 AI Agent 工作台。它支持聊天、流式响应、异步任务、模型切换、工具调用、Skills 扩展、长期记忆、向量检索、任务系统、webhook 和运行时控制台。",
        "我做它不是为了简单套一个大模型 API，而是想解决 Agent 产品真正上线时会遇到的问题：上下文怎么构建、长期记忆怎么管理、Function Calling 怎么保证安全、工具怎么扩展、运行过程怎么追踪、模型成本怎么统计。",
        "架构上，我把一次用户请求变成一个可追踪的 run。它会经过路由、上下文构建、模型决策、工具提案、权限和风险校验、执行、事件记录和记忆抽取。短期记忆用 Redis，长期记忆用 MySQL 的 memory_record，向量库用 Spring AI VectorStore 做语义召回增强，RabbitMQ 承接异步任务。",
        "其中我重点做的是三块：第一是 Agent Runtime，让不同入口统一进入一条可追踪链路；第二是工具和 Skills 治理，让模型能调用能力但不能越权执行；第三是记忆系统，让 Agent 能跨 session 记住用户偏好，同时避免乱记和错记。",
        "这个项目大部分代码实现和重构我都有用 Codex 辅助，但 Codex 主要是提高开发效率。真正的架构判断、边界设计、风险控制、测试验收和产品体验验证，是我自己负责的。"
    ]:
        add_body(doc, text)


def main():
    doc = Document(BASE)
    for p in doc.paragraphs:
        normalize_paragraph(p)
    replace_springclaw_resume_block(doc)
    append_interview_script(doc)

    section = doc.sections[-1]
    footer_p = section.footer.paragraphs[0]
    footer_p.text = "韩秉政｜Agent 开发面试材料"
    normalize_paragraph(footer_p, size=9)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
